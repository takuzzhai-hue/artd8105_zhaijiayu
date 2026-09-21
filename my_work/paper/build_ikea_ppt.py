#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "_local" / "py_pkgs"))

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from lxml import etree

ROOT = Path(__file__).resolve().parent
IMG = ROOT / "ikea_images"
OUT = ROOT / "從盒子到生活空間.pptx"

BG = RGBColor(0xF6, 0xF1, 0xE8)
INK = RGBColor(0x2A, 0x24, 0x1A)
MUTED = RGBColor(0x6B, 0x5E, 0x4E)
ACCENT = RGBColor(0xA8, 0x6F, 0x35)
RULE = RGBColor(0xD8, 0xC9, 0xB0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xEE, 0xE6, 0xD8)

FONT = "PingFang SC"
FONT_EN = "PingFang SC"

W, H = Inches(13.333), Inches(7.5)


def set_run(run, size, color, bold=False, font=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    # east asian font
    for child in list(rPr):
        if "ea" in child.tag:
            rPr.remove(child)
    ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", FONT)


def add_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_tb(slide, l, t, w, h, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=1.15):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.line_spacing = space
    run = p.add_run()
    run.text = text
    set_run(run, size, color, bold)
    return box


def add_paras(slide, l, t, w, h, items, size=16, color=INK, space=1.2, after=8):
    """items: list of (text, bold, color, size) or str"""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, str):
            text, b, c, s = item, False, color, size
        else:
            text = item[0]
            b = item[1] if len(item) > 1 else False
            c = item[2] if len(item) > 2 else color
            s = item[3] if len(item) > 3 else size
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = space
        p.space_after = Pt(after)
        run = p.add_run()
        run.text = text
        set_run(run, s, c, b)
    return box


def footer(slide, n, total, section=""):
    add_rect(slide, Inches(0), Inches(7.38), W, Inches(0.12), RULE)
    add_tb(slide, Inches(0.55), Inches(7.12), Inches(8), Inches(0.28),
           section, 10, MUTED)
    add_tb(slide, Inches(11.4), Inches(7.12), Inches(1.4), Inches(0.28),
           f"{n}  /  {total}", 10, MUTED, align=PP_ALIGN.RIGHT)


def kicker(slide, text):
    add_tb(slide, Inches(0.7), Inches(0.32), Inches(10), Inches(0.32),
           text, 12, ACCENT, bold=True)


def heading(slide, text, y=0.58, size=28):
    add_tb(slide, Inches(0.7), Inches(y), Inches(12), Inches(0.7),
           text, size, INK, bold=True)


def rule(slide, y=1.22):
    add_rect(slide, Inches(0.7), Inches(y), Inches(1.35), Inches(0.04), ACCENT)


def product_slide(prs, n, total, kicker_t, title, body_items, img_path, caption):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_rect(s, Inches(8.55), Inches(0), Inches(4.783), H, PANEL)
    kicker(s, kicker_t)
    heading(s, title, 0.58, 24)
    rule(s, 1.22)
    add_paras(s, Inches(0.7), Inches(1.45), Inches(7.5), Inches(5.4),
              body_items, size=16, after=10, space=1.18)
    # image card
    add_rect(s, Inches(8.85), Inches(1.05), Inches(4.15), Inches(4.15), WHITE)
    s.shapes.add_picture(str(img_path), Inches(8.95), Inches(1.15), Inches(3.95), Inches(3.95))
    add_tb(s, Inches(8.85), Inches(5.3), Inches(4.15), Inches(0.7),
           caption, 11, MUTED, align=PP_ALIGN.CENTER)
    footer(s, n, total, "二、五個 IKEA 紙製品")
    return s


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    total = 17

    # 1 title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_rect(s, Inches(0), Inches(0), Inches(0.22), H, ACCENT)
    add_tb(s, Inches(0.9), Inches(2.05), Inches(11.5), Inches(0.4),
           "從盒子到生活空間", 14, ACCENT, bold=True)
    add_tb(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.4),
           "從盒子到生活空間", 44, INK, bold=True)
    add_tb(s, Inches(0.9), Inches(4.05), Inches(11.5), Inches(0.7),
           "IKEA 紙製收納與臨時性的設計", 24, MUTED)
    add_rect(s, Inches(0.9), Inches(4.9), Inches(1.5), Inches(0.05), ACCENT)
    add_tb(s, Inches(0.9), Inches(5.15), Inches(11), Inches(0.4),
           "From the box to living space  ·  IKEA paper storage and the design of the temporary",
           13, MUTED)
    footer(s, 1, total, "")

    # 2 intro
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    kicker(s, "一")
    heading(s, "從 IKEA 的紙盒開始")
    rule(s)
    add_paras(s, Inches(0.7), Inches(1.55), Inches(12), Inches(5.2), [
        ("在日本 IKEA 的收納區，紙與紙板被大量使用；但在國內的收納區，紙製收納的比例明顯低得多。", False, INK, 20),
        ("我們通常認為收納用品應該堅固、防水、耐久，但 IKEA 卻持續使用紙板這種相對脆弱的材料。", False, INK, 20),
    ], after=18, space=1.35)
    footer(s, 2, total, "一、從 IKEA 的紙盒開始")

    # 3 paper qualities
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    kicker(s, "一")
    heading(s, "紙板具有一些與傳統家具完全不同的特點")
    rule(s)
    words = ["輕", "可折疊", "可組裝", "可替換", "可回收"]
    x0 = 0.7
    for i, w in enumerate(words):
        left = Inches(x0 + i * 2.45)
        add_rect(s, left, Inches(2.2), Inches(2.2), Inches(2.35), WHITE)
        add_rect(s, left, Inches(2.2), Inches(2.2), Inches(0.08), ACCENT)
        add_tb(s, left, Inches(2.95), Inches(2.2), Inches(0.9),
               w, 28, INK, bold=True, align=PP_ALIGN.CENTER)
    footer(s, 3, total, "一、從 IKEA 的紙盒開始")

    # 4 question
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_rect(s, Inches(0), Inches(0), Inches(0.22), H, ACCENT)
    add_tb(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(0.4),
           "於是我開始思考", 14, ACCENT, bold=True)
    add_tb(s, Inches(0.9), Inches(2.95), Inches(11.5), Inches(2.2),
           "為什麼一種看起來並不永久的材料，\n可以長期參與我們的日常生活與空間組織？",
           28, INK, bold=True)
    footer(s, 4, total, "一、從 IKEA 的紙盒開始")

    # 5 five products overview
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    kicker(s, "二")
    heading(s, "五個 IKEA 紙製品")
    rule(s)
    add_tb(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.55),
           "這五個產品完成尺度的升級：從一張紙，到容器，到系統，到家具。", 18, MUTED)
    steps = [
        ("①", "FJÄDERHARV 收納盒", "紙 → 容器"),
        ("②", "FJÄDERHARV 雜誌文件盒", "容器 → 系統"),
        ("③", "FJÄDERHARV 迷你抽屜櫃", "容器 → 家具"),
        ("④", "GÄRDESGÅRD", "公開材料"),
        ("⑤", "KVARNVIK", "掩飾材料"),
    ]
    for i, (num, name, tag) in enumerate(steps):
        top = Inches(2.15 + i * 0.9)
        add_rect(s, Inches(0.7), top, Inches(11.9), Inches(0.78), WHITE)
        add_tb(s, Inches(0.9), top + Inches(0.16), Inches(0.7), Inches(0.5),
               num, 20, ACCENT, bold=True)
        add_tb(s, Inches(1.7), top + Inches(0.16), Inches(7), Inches(0.5),
               name, 20, INK, bold=True)
        add_tb(s, Inches(8.7), top + Inches(0.16), Inches(3.6), Inches(0.5),
               tag, 18, MUTED, align=PP_ALIGN.RIGHT)
    footer(s, 5, total, "二、五個 IKEA 紙製品")

    # 6 product 1
    product_slide(
        prs, 6, total, "二  ·  ①",
        "FJÄDERHARV 收納盒  ——  紙 → 容器",
        [
            ("一張平面紙板經過折疊與組裝，形成六個面，並產生一個內部空間。", False, INK, 18),
            ("平板包裝時體積幾乎為零，空間是在使用者手上生成。", False, INK, 18),
        ],
        IMG / "fjaederharv_box.jpg",
        "FJÄDERHARV 收納盒\n來源：IKEA 官網產品圖",
    )

    # 7 product 2
    product_slide(
        prs, 7, total, "二  ·  ②",
        "FJÄDERHARV 雜誌文件盒  ——  容器 → 系統",
        [
            ("雜誌盒的尺寸並不是獨立存在的，它必須對應一條尺寸鏈：", False, INK, 16),
            ("文件 → 收納盒 → 書架 → 家具 → 房間", True, ACCENT, 16),
            ("它的價值不在自身強度，而在能否嵌入既有的尺寸系統。", False, INK, 16),
            ("斜切的開口：用「減少材料」換取可視性與抽取的動作。", False, INK, 16),
            ("單面開放：放棄密閉性，換來單手操作。", False, INK, 16),
            ("一個很小的紙製品，進入了更大的空間系統，參與了生活空間的組織。", False, INK, 16),
        ],
        IMG / "fjaederharv_magazine.jpg",
        "FJÄDERHARV 雜誌文件盒\n來源：IKEA 官網產品圖",
    )

    # 8 product 3
    product_slide(
        prs, 8, total, "二  ·  ③",
        "FJÄDERHARV 迷你抽屜櫃  ——  容器 → 家具",
        [
            ("它仍然主要由紙板構成，但增加了：外框 / 抽屜 / 抽拉結構。", False, INK, 16),
            ("外框承重、抽屜收納，把紙不能受力的弱點交給框架，紙只負責圍合。", False, INK, 16),
            ("抽屜引入了「動作」：拉出、推入。", False, INK, 16),
            ("多個小抽屜意味著分類，它從「放東西的地方」變成「產生秩序的裝置」。", False, INK, 16),
        ],
        IMG / "fjaederharv_chest.jpg",
        "FJÄDERHARV 迷你抽屜櫃\n來源：IKEA 官網產品圖",
    )

    # 9 product 4
    product_slide(
        prs, 9, total, "二  ·  ④",
        "GÄRDESGÅRD  ——  公開材料",
        [
            ("與白色的 FJÄDERHARV 不同，GÄRDESGÅRD 保留了非常明顯的原色紙板質感與瓦楞紋理。", False, INK, 16),
            ("它沒有試圖隱藏自己是一種普通、廉價的紙材料，材料的身份是公開的。", False, INK, 16),
            ("引入「可替換」的概念，誠實的材料表達，降低了使用者的心理成本。", False, INK, 16),
        ],
        IMG / "gaerdesgard.jpg",
        "GÄRDESGÅRD\n來源：IKEA 官網產品圖",
    )

    # 10 product 5
    product_slide(
        prs, 10, total, "二  ·  ⑤",
        "KVARNVIK  ——  掩飾材料",
        [
            ("KVARNVIK 同樣以紙板為主要結構，但透過表面處理、把手、邊緣包邊與蓋子的合縫，看起來更像一件可以長期保存物品的正式收納家具。", False, INK, 16),
            ("它用細節借來了「耐久家具」的語言。", False, INK, 16),
            ("材料是臨時的，但設計讓我們產生永久的期待。", False, INK, 16),
        ],
        IMG / "kvarnvik.jpg",
        "KVARNVIK\n來源：IKEA 官網產品圖",
    )

    # 11 why paper
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    kicker(s, "三")
    heading(s, "為什麼是紙？")
    rule(s)
    add_tb(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5),
           "這五個產品共同具有一些特點：", 18, MUTED)
    words = ["輕", "可折疊", "可替換", "可調整", "非永久"]
    for i, w in enumerate(words):
        left = Inches(0.7 + i * 2.45)
        add_rect(s, left, Inches(2.2), Inches(2.2), Inches(1.7), WHITE)
        add_tb(s, left, Inches(2.7), Inches(2.2), Inches(0.8),
               w, 26, INK, bold=True, align=PP_ALIGN.CENTER)
    add_tb(s, Inches(0.7), Inches(4.3), Inches(12), Inches(1.4),
           "通常我們會把這些特徵理解成材料的缺點。", 22, INK, bold=True)
    footer(s, 11, total, "三、為什麼是紙？")

    # 12 Japanese paper
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    kicker(s, "四")
    heading(s, "日本傳統空間中的紙")
    rule(s)
    add_tb(s, Inches(0.7), Inches(1.45), Inches(12), Inches(0.5),
           "這讓我想到日本傳統生活中紙的使用：", 16, MUTED)
    items = [
        ("障子", "用紙形成透光的空間邊界"),
        ("襖", "可以移動，改變房間之間的關係"),
        ("和紙包裝", "暫時包裹與保護物體"),
        ("折形", "透過折疊產生邊界與容器"),
    ]
    for i, (t, d) in enumerate(items):
        col = i % 2
        row = i // 2
        left = Inches(0.7 + col * 6.2)
        top = Inches(2.05 + row * 1.55)
        add_rect(s, left, top, Inches(5.85), Inches(1.38), WHITE)
        add_rect(s, left, top, Inches(0.1), Inches(1.38), ACCENT)
        add_tb(s, left + Inches(0.4), top + Inches(0.22), Inches(5.2), Inches(0.45),
               t, 20, INK, bold=True)
        add_tb(s, left + Inches(0.4), top + Inches(0.7), Inches(5.2), Inches(0.5),
               d, 16, MUTED)
    add_tb(s, Inches(0.7), Inches(5.3), Inches(12), Inches(1.5),
           "紙並不是最堅固的材料，卻長期承擔著重要的空間功能。\n它們有一個共同特點：不追求用最堅固的材料永久固定空間，而是允許空間被改變、替換與重新組織。",
           16, INK)
    footer(s, 12, total, "四、日本傳統空間中的紙")

    # 13 fragility as function
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    kicker(s, "五")
    heading(s, "脆弱也可以是一種功能")
    rule(s)
    add_paras(s, Inches(0.7), Inches(1.5), Inches(12), Inches(5.3), [
        ("我們通常認為好的設計應該是：耐用、堅固、長久。", False, INK, 18),
        ("但在實際使用 IKEA 紙盒的過程中，我發現它恰恰不是這樣：不能清洗、污漬很難擦除、容易損壞、不能承受太重的物品。", False, INK, 18),
        ("但同時，紙也因此獲得另一種特性：輕、可替換、可折疊、非永久。", False, INK, 18),
        ("障子的紙損壞以後，可以重新替換。紙盒也一樣——方便打包帶走，方便斷捨離，方便經歷每一段生活。", False, INK, 18),
    ], after=14, space=1.25)
    footer(s, 13, total, "五、脆弱也可以是一種功能")

    # 14 performance of change
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_rect(s, Inches(0), Inches(0), Inches(0.22), H, ACCENT)
    add_tb(s, Inches(0.9), Inches(2.55), Inches(11.5), Inches(2.2),
           "如果人的生活本身不斷變化，那麼容易移動、替換與重新組織，也可能是一種性能。\n\n而在我的生活中，我確實不斷需要紙盒來重新建立秩序。",
           24, INK, bold=True)
    footer(s, 14, total, "五、脆弱也可以是一種功能")

    # 15 Imoto
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    kicker(s, "六")
    heading(s, "井本佐保里的居住研究")
    rule(s)
    add_paras(s, Inches(0.7), Inches(1.5), Inches(12), Inches(5.3), [
        ("這讓我聯想到井本佐保里關於居住環境、災後臨時居住、住宅重建與社區生活的研究。", False, INK, 18),
        ("她關注的一個重要問題是：人在不斷變化的環境中，如何繼續生活。", True, INK, 18),
        ("在災後的居住環境中，人的生活並不是簡單地：", False, INK, 18),
        ("舊住宅  →  新住宅", True, ACCENT, 20),
        ("而是可能經歷：", False, INK, 18),
        ("臨時居住  →  移動  →  調整  →  關係重新形成  →  新的居住環境", True, ACCENT, 18),
        ("從這個角度看，「臨時」並不意味著生活停止了。即使空間是臨時的、物品是臨時的，人仍然需要整理物品、建立關係、形成秩序，並繼續生活。", False, INK, 17),
    ], after=10, space=1.2)
    footer(s, 15, total, "六、井本佐保里的居住研究")

    # 16 reread boxes
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    kicker(s, "七")
    heading(s, "重新理解 IKEA 的紙盒")
    rule(s)
    add_paras(s, Inches(0.7), Inches(1.5), Inches(12), Inches(5.3), [
        ("再回來看這五個紙製品，我不再把它們理解為「不夠耐用的收納家具」。", False, INK, 18),
        ("恰恰相反：臨時性的物件，也可以支持持續的生活。", True, INK, 20),
        ("紙盒會損壞、會被替換、會隨著搬家與生活階段而消失。", False, INK, 18),
        ("但「整理生活」這件事不會消失。每當生活發生變化，我們又需要重新尋找新的容器、新的位置與新的秩序。", False, INK, 18),
    ], after=16, space=1.28)
    footer(s, 16, total, "七、重新理解 IKEA 的紙盒")

    # 17 closing
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s)
    add_rect(s, Inches(0), Inches(0), Inches(0.22), H, ACCENT)
    add_tb(s, Inches(0.9), Inches(1.35), Inches(11.5), Inches(0.4),
           "因此，我最後想提出的問題是", 14, ACCENT, bold=True)
    add_tb(s, Inches(0.9), Inches(1.85), Inches(11.5), Inches(1.4),
           "如果生活本身一直在變化，為什麼設計一定要永久？",
           28, INK, bold=True)
    add_rect(s, Inches(0.9), Inches(3.4), Inches(1.4), Inches(0.05), ACCENT)
    add_tb(s, Inches(0.9), Inches(3.7), Inches(11.5), Inches(2.8),
           "也許接受紙的「脆弱」：接受物品不會永遠存在，也接受生活本身具有脆弱性與不確定性。\n\n設計不一定需要對抗這種變化。它也可以幫助我們在變化發生之後，再一次建立生活。",
           18, MUTED)
    footer(s, 17, total, "七、重新理解 IKEA 的紙盒")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
